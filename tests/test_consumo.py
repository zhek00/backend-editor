"""Registro central do consumo no OpenRouter (fabrica/consumo.py): o que entra, o que fica de fora e a planilha."""
import json

import pytest

from fabrica import consumo


@pytest.fixture
def pasta(tmp_path, monkeypatch):
    monkeypatch.setattr(consumo, "PASTA", tmp_path / "relatorios")
    monkeypatch.setattr(consumo, "REGISTRO", tmp_path / "relatorios" / "openrouter_consumo.csv")
    monkeypatch.setattr(consumo, "PLANILHA", tmp_path / "relatorios" / "consumo_openrouter.xlsx")
    monkeypatch.setattr(consumo, "OFICIAL", tmp_path / "relatorios" / "oficial.json")
    monkeypatch.setattr(consumo, "RAIZ", tmp_path)
    monkeypatch.setattr(consumo, "agendar", lambda: None)
    monkeypatch.setattr(consumo, "_oficial", lambda: {"chave": {"usage": 1, "usage_monthly": 1, "usage_weekly": 1,
                                                                "usage_daily": 0, "limit": 10, "limit_remaining": 9},
                                                      "creditos": {"total_credits": 30}, "consulta": "teste"})
    return tmp_path


class _P:
    nome = "video-1"


def test_cada_chamada_vira_uma_linha(pasta):
    consumo.registrar(_P(), "texto", "escolha de material real", "qwen/qwen3.8-flash", 1000, 200, 0.0017)
    consumo.registrar(_P(), "texto", "julgar mídia", "dots-studio/dots-3-note-preview:free", 900, 50, 0)
    linhas = consumo.ler()
    assert [l["modelo"] for l in linhas] == ["qwen/qwen3.8-flash", "dots-studio/dots-3-note-preview:free"]
    assert [l["gratuito"] for l in linhas] == ["Não", "Sim"]
    assert float(linhas[0]["custo_usd"]) == 0.0017 and linhas[0]["projeto"] == "video-1"


def test_so_o_que_passou_pelo_openrouter(pasta):
    consumo.registrar_custo(_P(), "imagem", {"provedor": "openrouter", "modelo": "openai/gpt-5.4-image-2"}, 0.0046)
    consumo.registrar_custo(_P(), "imagem", {"provedor": "kie", "modelo": "grok"}, 0.02)
    consumo.registrar_custo(_P(), "narracao", {"provedor": "genaipro"}, 0.05)
    assert [l["tipo"] for l in consumo.ler()] == ["imagem"]


def test_historico_entra_uma_vez(pasta):
    projeto = pasta / "projetos" / "antigo"
    projeto.mkdir(parents=True)
    (projeto / "uso_openrouter.json").write_text(json.dumps([
        {"quando": "2026-10-05T10:00:00", "etapa": "diretor", "modelo": "qwen/qwen3.8-flash", "tokens_lidos": 10,
         "tokens_escritos": 5, "custo_usd": 0.01},
        {"quando": "2026-10-05T10:00:01", "etapa": "x", "modelo": "aimlapi:outro", "tokens_lidos": 1,
         "tokens_escritos": 1, "custo_usd": 0.5}]), encoding="utf-8")
    assert consumo.importar_historico(log=lambda *a: None) == 1  # a AIMLAPI é outra conta
    assert consumo.importar_historico(log=lambda *a: None) == 0  # rodar de novo não repete


def test_planilha_sai_e_espera_quando_esta_aberta(pasta, monkeypatch):
    consumo.registrar(_P(), "texto", "diretor", "qwen/qwen3.8-flash", 10, 5, 0.01)
    assert consumo.gerar_planilha(log=lambda *a: None) is True
    assert consumo.PLANILHA.exists()

    def aberta(*a):
        raise PermissionError("aberta no Excel")
    monkeypatch.setattr(consumo.os, "replace", aberta)
    assert consumo.gerar_planilha(log=lambda *a: None) is False
    assert consumo._pendente is True  # tenta de novo depois


def test_csv_para_o_google_e_a_chave_propria(pasta, monkeypatch):
    monkeypatch.setenv("CONSUMO_PLANILHA_TOKEN", "chave-de-leitura")
    consumo.registrar(_P(), "texto", "diretor", "qwen/qwen3.8-flash", 10, 5, 0.0017)
    consumo.registrar(_P(), "texto", "julgar", "dots-studio/dots-3-note-preview:free", 10, 5, 0)
    linhas = consumo.csv_do_registro().splitlines()
    assert linhas[0] == ",".join(consumo.COLUNAS_CSV)
    data, hora, *_resto, paga, lidos, escritos, custo = linhas[1].split(",")
    assert len(data) == 10 and len(hora) == 5 and paga == "1" and custo == "0.0017"
    assert linhas[2].split(",")[7] == "0"  # gratuita: não conta como paga
    assert consumo.chave_confere("chave-de-leitura") and not consumo.chave_confere("outra")
    assert not consumo.chave_confere("")


def test_modelo_do_google_puxa_o_registro(pasta, monkeypatch):
    from openpyxl import load_workbook
    monkeypatch.setenv("CONSUMO_PLANILHA_TOKEN", "chave-de-leitura")
    arquivo = consumo.gerar_modelo_google("https://exemplo.test", pasta / "g.xlsx")
    wb = load_workbook(arquivo)
    assert wb["Chamadas"]["A1"].value.startswith('=IMPORTDATA("https://exemplo.test/consumo/openrouter.csv?chave=')
    assert wb["Custo diário"]["A4"].value.startswith("=QUERY(Chamadas!A:K")
    assert set(wb.sheetnames) >= {"Resumo", "Chamadas", "Custo diário", "Custo por API", "Diário por API",
                                  "Por etapa", "Por projeto"}
